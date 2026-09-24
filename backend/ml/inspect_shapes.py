import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation(r'Презентация/Презентация/ЛЦТ2026 Шаблон презентации.pptx')
for idx in [26, 27, 28]:
    s = prs.slides[idx]
    print(f'=== SLIDE {idx+1} ===')
    for sh in s.shapes:
        txt = sh.text.strip().replace('\n', ' ') if sh.has_text_frame else ''
        print(f'  name={sh.name}, type={sh.shape_type}, text="{txt[:50]}"')
