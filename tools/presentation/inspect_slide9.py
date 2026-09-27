import sys
from pathlib import Path
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
p = Path('presentation/Москоллектор_НейроКонтур_Защита.pptx')
prs = Presentation(str(p))
s9 = prs.slides[8]  # 0-indexed slide 9
print(f"=== Slide 9 Shapes (Total: {len(s9.shapes)}) ===")
for i, sh in enumerate(s9.shapes):
    t = sh.text_frame.text if sh.has_text_frame else ''
    t_clean = ' '.join(t.split())
    print(f"Shape {i}: '{sh.name}' ({sh.shape_type}), left={sh.left}, top={sh.top}, text='{t_clean[:120]}'")
