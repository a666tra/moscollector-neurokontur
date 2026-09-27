import sys
from pathlib import Path
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
p = Path('presentation/Москоллектор_НейроКонтур_Защита.pptx')
prs = Presentation(str(p))

for slide_idx in [10, 13]:  # Slide 11 (idx 10), Slide 14 (idx 13)
    s = prs.slides[slide_idx]
    print(f"\n=== Slide {slide_idx+1} Shapes (Total: {len(s.shapes)}) ===")
    for i, sh in enumerate(s.shapes):
        t = sh.text_frame.text if sh.has_text_frame else ''
        t_clean = ' '.join(t.split())
        if t_clean:
            print(f"Shape {i} ('{sh.name}'): '{t_clean}'")
