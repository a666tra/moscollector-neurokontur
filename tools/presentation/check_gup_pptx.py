import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation('presentation/Москоллектор_НейроКонтур_Защита.pptx')
found = False
for s_idx, slide in enumerate(prs.slides):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if 'ГУП' in p.text or 'GUP' in p.text:
                    print(f"Slide {s_idx+1}: {p.text}")
                    found = True
if not found:
    print("No GUP found in presentation PPTX.")
