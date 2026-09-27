import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation('presentation/Москоллектор_НейроКонтур_Защита.pptx')
for idx, slide in enumerate(prs.slides):
    print(f'=== SLIDE {idx+1} TOP SHAPES ===')
    for s in slide.shapes:
        top_in = s.top / 914400
        if top_in < 1.5 and not 'Номер слайда' in s.name:
            txt = s.text_frame.text.replace('\n', ' ') if s.has_text_frame else '<no text>'
            print(f'  [{s.name}] ({s.left/914400:.2f},{s.top/914400:.2f} {s.width/914400:.2f}x{s.height/914400:.2f}) text="{txt}"')
