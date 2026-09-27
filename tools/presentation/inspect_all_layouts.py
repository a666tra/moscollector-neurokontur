import sys
from pathlib import Path
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
p = Path('presentation').glob('*.bak')
bak_file = list(p)[0]
prs = Presentation(str(bak_file))
s1 = prs.slides[0]

print("=== Slide Master Shapes ===")
for sm in prs.slide_masters:
    for sh in sm.shapes:
        t = sh.text_frame.text if sh.has_text_frame else ''
        print(f"Master Shape: '{sh.name}', type={sh.shape_type}, text='{t.strip()}'")

print("=== Layouts ===")
for i, l in enumerate(prs.slide_layouts):
    for sh in l.shapes:
        t = sh.text_frame.text if sh.has_text_frame else ''
        if t or sh.shape_type != 1:
            print(f"Layout {i} ('{l.name}'): '{sh.name}', type={sh.shape_type}, text='{t.strip()}'")
