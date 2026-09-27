import sys
import pptx

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation(r'Презентация/Презентация/ЛЦТ2026 Шаблон презентации.pptx')

for idx in [7, 8, 9, 10]:  # 7=Slide 8, 8=Slide 9, 9=Slide 10, 10=Slide 11
    slide = prs.slides[idx]
    print(f"=== SLIDE {idx+1} ===")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            full_text = " | ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if full_text:
                print(f"  Shape {s_idx} ({shape.name}): {full_text[:120]}")
