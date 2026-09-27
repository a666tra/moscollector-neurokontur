import pymupdf
import subprocess
from pptx import Presentation
from pathlib import Path

prs = Presentation('presentation/Москоллектор_НейроКонтур_Защита.pptx.bak')
s1 = prs.slides[0]
s2 = prs.slides[1]
image2_part = s2.part.rels['rId2']._target
s1.part.rels['rId2']._target = image2_part

test_pptx = Path('presentation/test_s1.pptx')
test_pdf = Path('presentation/test_s1.pdf')
prs.save(str(test_pptx))
print("Saved test_s1.pptx")

# Export to PDF via PowerPoint COM
ps_cmd = f"""
$ppt = New-Object -ComObject PowerPoint.Application
$pres = $ppt.Presentations.Open('{test_pptx.resolve()}')
$pres.SaveAs('{test_pdf.resolve()}', 32)
$pres.Close()
$ppt.Quit()
"""
res = subprocess.run(['powershell', '-Command', ps_cmd], capture_output=True, text=True)
print("PowerPoint export output:", res.stdout, res.stderr)

if test_pdf.exists():
    doc = pymupdf.open(str(test_pdf))
    page = doc[0]
    pix = page.get_pixmap(dpi=150)
    pix.save('presentation/test_s1_rendered.png')
    print("Saved presentation/test_s1_rendered.png")
