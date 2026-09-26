import pptx

prs = pptx.Presentation('presentation/Москоллектор_НейроКонтур_Защита.pptx')
with open('presentation_dump.txt', 'w', encoding='utf-8') as out:
    for idx, slide in enumerate(prs.slides):
        out.write(f"=== SLIDE {idx+1} ===\n")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    txt = p.text.strip()
                    if txt:
                        out.write(f"  {txt}\n")
            elif shape.has_table:
                out.write("  [TABLE]:\n")
                for row in shape.table.rows:
                    row_txt = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    out.write("    | " + " | ".join(row_txt) + " |\n")
print("Dumped successfully")
