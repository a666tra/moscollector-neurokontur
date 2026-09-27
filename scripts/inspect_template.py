import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation('Шаблон2026 Презентация проекта.pptx')
s1 = prs.slides[0]
print(f"Template slide 1 shapes ({len(s1.shapes)}):")
for sh in s1.shapes:
    t = sh.text_frame.text if sh.has_text_frame else ''
    print(f"  Shape '{sh.name}': {t.strip()}")
