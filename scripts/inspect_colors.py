import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation('presentation/Москоллектор_НейроКонтур_Защита.pptx')
for idx in [3, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17]:
    slide = prs.slides[idx]
    print(f'=== SLIDE {idx+1} TEXT RUN COLORS ===')
    for s_idx, s in enumerate(slide.shapes):
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        c_type = r.font.color.type if r.font.color else 'no_color'
                        rgb = str(r.font.color.rgb) if r.font.color and r.font.color.type == 1 else 'not_rgb'
                        theme_color = str(r.font.color.theme_color) if r.font.color and hasattr(r.font.color, 'theme_color') else 'none'
                        print(f'  Shape #{s_idx} ({s.name}) text="{r.text[:30]}..." color_type={c_type} rgb={rgb} theme={theme_color}')
                        break
