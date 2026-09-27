import sys
from pathlib import Path
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
p = list(Path('presentation').glob('*.bak'))[0]
prs = Presentation(str(p))

for idx in [0, 1, 2, 7, 9, 22]:
    l = prs.slide_layouts[idx]
    print(f"\n--- Layout {idx}: '{l.name}' ---")
    for s in l.shapes:
        t = s.text_frame.text if s.has_text_frame else ''
        print(f"  Shape: '{s.name}' ({s.shape_type}), text='{t.strip()}'")
