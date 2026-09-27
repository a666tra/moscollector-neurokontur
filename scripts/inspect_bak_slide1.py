import sys
from pathlib import Path
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
p = Path('presentation').glob('*.bak')
bak_file = list(p)[0]
print(f"Loading {bak_file}")
prs = Presentation(str(bak_file))
s1 = prs.slides[0]
print(f"Slide 1 has {len(s1.shapes)} shapes:")
for i, sh in enumerate(s1.shapes):
    t = sh.text_frame.text if sh.has_text_frame else ''
    print(f"  Shape {i}: id={sh.shape_id}, name='{sh.name}', text='{t.strip()}', left={sh.left}, top={sh.top}, w={sh.width}, h={sh.height}")
