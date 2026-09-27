import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation('presentation/Москоллектор_НейроКонтур_Защита.pptx')
s1 = prs.slides[0]
print(f"Slide 1 Layout: '{s1.slide_layout.name}'")
print("=== Slide 1 Shapes ===")
for i, sh in enumerate(s1.shapes):
    txt = sh.text_frame.text if sh.has_text_frame else ''
    print(f"Slide Shape {i}: '{sh.name}', text='{txt.strip()}'")

print("=== Slide 1 Layout Shapes ===")
for i, sh in enumerate(s1.slide_layout.shapes):
    txt = sh.text_frame.text if sh.has_text_frame else ''
    print(f"Layout Shape {i}: '{sh.name}', text='{txt.strip()}'")

print("=== Slide Master Shapes ===")
for i, sh in enumerate(s1.slide_layout.slide_master.shapes):
    txt = sh.text_frame.text if sh.has_text_frame else ''
    print(f"Master Shape {i}: '{sh.name}', text='{txt.strip()}'")
